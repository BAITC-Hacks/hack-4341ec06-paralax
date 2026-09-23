"""Stage 5: extract source balances without inventing a current stock snapshot.

The partner workbooks are read from the immutable source Git commit. Derived
values are written to a local ignored JSONL file; no workbook is modified.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
import re
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from clean_demand import SOURCE_COMMIT

STAGE1_FILE = Path(".analysis/clean-demand/all-skus.jsonl")
OUTPUT_FILE = Path(".analysis/balances/source-snapshot.jsonl")
IEK_TRANSIT_FILE = "data/IEK/Путь ИЭК 22.09.2026.xlsx"
SYSTEME_SNAPSHOT_FILE = "data/Systeme electric/Товар в пути_SystemElectric на 22.09.2026.xlsx"
IEK_SHEET = "Лист4"
SYSTEME_SHEET = "TDSheet"
IEK_SNAPSHOT_DATE_HINT = "2026-09-22"
SYSTEME_DATE_NOTE = (
    "Единая дата и охват сводки не подтверждены: имя файла указывает 22.09.2026, "
    "а заголовок товара в пути — 24.09; смысл второй даты требует уточнения."
)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_git_sheet(path: str, name: str):
    source = subprocess.check_output(["git", "show", f"{SOURCE_COMMIT}:{path}"])
    workbook = load_workbook(io.BytesIO(source), read_only=True, data_only=True)
    try:
        sheet = workbook[name]
    except KeyError:
        workbook.close()
        raise
    return workbook, sheet


def numeric_quantity(value: object) -> int | float | None:
    """Accept only explicit finite Excel numbers; blanks and numeric text differ."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return value if math.isfinite(value) and value >= 0 else None


def blank(value: object) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def cell_ref(path: str, sheet: str, row: int, column: int, **details: object) -> dict:
    return {
        "file": path,
        "sheet": sheet,
        "row": row,
        "column": column,
        "cell": f"{get_column_letter(column)}{row}",
        **details,
    }


def evidence(
    quantity: int | float | None,
    quality: str,
    as_of_date: str | None,
    basis: str,
    source_refs: list[dict],
    scope_note: str,
) -> dict:
    return {
        "quantity": quantity,
        "quality": quality,
        "as_of_date": as_of_date,
        "basis": basis,
        "source_refs": source_refs,
        "scope_note": scope_note,
    }


def shipment_headers(row: Sequence[object]) -> list[dict]:
    if not row or str(row[0]).strip() != "Код 1с":
        raise ValueError("Unexpected IEK transit header in A1")
    result = []
    for column in range(4, 10):
        title = str(row[column - 1]) if len(row) >= column and row[column - 1] is not None else ""
        shipment = re.search(r"\bУТ-\d+\b", title)
        arrival = re.search(r"поступление\s+до\s+(\d{2}\.\d{2}\.\d{4})", title, re.IGNORECASE)
        if not shipment or not arrival:
            raise ValueError(f"Unrecognized IEK shipment header in {get_column_letter(column)}1: {title!r}")
        try:
            latest = datetime.strptime(arrival.group(1), "%d.%m.%Y").date().isoformat()
        except ValueError as error:
            raise ValueError(f"Invalid IEK arrival date in {get_column_letter(column)}1") from error
        result.append({"column": column, "shipment_id": shipment.group(),
                       "expected_arrival_latest": latest, "shipment_label": title})
    return result


def parse_iek_transit_rows(
    header: Sequence[object], numbered_rows: Iterable[tuple[int, Sequence[object]]]
) -> dict[str, dict]:
    """Collect all shipment cells, including blanks and repeated SKU rows."""
    shipments = shipment_headers(header)
    collected: dict[str, dict] = {}
    for row_number, row in numbered_rows:
        raw_sku = row[0] if row else None
        if raw_sku is None or not str(raw_sku).strip():
            continue
        sku = str(raw_sku).strip()
        entry = collected.setdefault(sku, {"amounts": [], "source_refs": [], "notes": [], "rows": []})
        entry["rows"].append(row_number)
        for shipment in shipments:
            column = shipment["column"]
            raw = row[column - 1] if len(row) >= column else None
            entry["source_refs"].append(cell_ref(
                IEK_TRANSIT_FILE, IEK_SHEET, row_number, column,
                shipment_id=shipment["shipment_id"],
                expected_arrival_latest=shipment["expected_arrival_latest"],
                arrival_basis="on_or_before",
                shipment_label=shipment["shipment_label"],
                source_value=raw if isinstance(raw, (str, int, float)) else None,
            ))
            if blank(raw):
                continue
            amount = numeric_quantity(raw)
            if amount is None:
                entry["notes"].append(f"invalid_transit_cell:{get_column_letter(column)}{row_number}")
            else:
                entry["amounts"].append(amount)
    result = {}
    for sku, entry in collected.items():
        if len(entry["rows"]) > 1:
            entry["notes"].append("duplicate_sku_rows_combined_by_shipment_cells")
        if any(note.startswith("invalid_transit_cell:") for note in entry["notes"]):
            quantity, quality = None, "invalid_source"
        elif not entry["amounts"]:
            quantity, quality = None, "missing"
            entry["notes"].append("all_shipment_cells_blank")
        else:
            quantity, quality = math.fsum(entry["amounts"]), "ambiguous_date"
        result[sku] = {
            "evidence": evidence(
                quantity, quality, IEK_SNAPSHOT_DATE_HINT, "shipment_columns_D_to_I",
                entry["source_refs"],
                "Дата 22.09.2026 взята из имени файла; полнота партий и отсутствие пересечения "
                "с текущим остатком не подтверждены. ETA означает 'поступление до', не точный день прихода."
            ),
            "notes": entry["notes"],
        }
    return result


def parse_systeme_rows(numbered_rows: Iterable[tuple[int, Sequence[object]]]) -> dict[str, dict]:
    """Read explicit AX/BC numbers; duplicate rows are conflicting evidence."""
    collected: dict[str, list[tuple[int, Sequence[object]]]] = {}
    for row_number, row in numbered_rows:
        raw_sku = row[2] if len(row) > 2 else None
        if raw_sku is None or not str(raw_sku).strip():
            continue
        collected.setdefault(str(raw_sku).strip(), []).append((row_number, row))

    result = {}
    for sku, entries in collected.items():
        notes = []
        if len(entries) > 1:
            notes.append("duplicate_sku_rows_conflicting_snapshot")
        fields = {}
        for kind, column, basis in (
            ("current_stock", 50, "reported_total_stock_AX"),
            ("goods_in_transit", 55, "reported_goods_in_transit_BC"),
        ):
            refs = []
            values = []
            for row_number, row in entries:
                raw = row[column - 1] if len(row) >= column else None
                refs.append(cell_ref(
                    SYSTEME_SNAPSHOT_FILE, SYSTEME_SHEET, row_number, column,
                    source_value=raw if isinstance(raw, (str, int, float)) else None,
                ))
                values.append(raw)
            if len(entries) > 1:
                quantity, quality = None, "conflict"
            else:
                quantity = numeric_quantity(values[0])
                quality = "ambiguous_date" if quantity is not None else (
                    "missing" if blank(values[0]) else "invalid_source"
                )
                if quality == "invalid_source":
                    notes.append(f"invalid_{kind}_cell:{refs[0]['cell']}")
            fields[kind] = evidence(quantity, quality, None, basis, refs, SYSTEME_DATE_NOTE)

        if len(entries) == 1:
            row_number, row = entries[0]
            stock = fields["current_stock"]["quantity"]
            warehouses = [numeric_quantity(row[column - 1]) if len(row) >= column else None
                          for column in range(46, 50)]  # AT:AW
            if stock is not None and all(value is not None for value in warehouses):
                total = math.fsum(warehouses)
                if not math.isclose(total, stock, abs_tol=1e-9):
                    notes.append(
                        f"warehouse_scope_differs:AT{row_number}:AW{row_number}="
                        f"{total:g},AX{row_number}={stock:g}"
                    )
        result[sku] = {**fields, "notes": notes}
    return result


def load_source_balances() -> tuple[dict[str, dict], dict[str, dict]]:
    workbook, sheet = read_git_sheet(IEK_TRANSIT_FILE, IEK_SHEET)
    try:
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        iek = parse_iek_transit_rows(
            header, enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2))
    finally:
        workbook.close()

    workbook, sheet = read_git_sheet(SYSTEME_SNAPSHOT_FILE, SYSTEME_SHEET)
    try:
        if sheet["C2"].value != "Код 1с" or sheet["AX2"].value != "Остаток" \
                or not str(sheet["BC2"].value).startswith("СЭ в пути"):
            raise ValueError("Unexpected Systeme snapshot headers in C2/AX2/BC2")
        systeme = parse_systeme_rows(
            enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3))
    finally:
        workbook.close()
    return iek, systeme


def iek_historical_stock(stage1: dict) -> tuple[dict, list[str]]:
    notes = []
    septembers = [period for period in stage1.get("periods", [])
                  if period.get("period") == "2026-09"]
    if len(septembers) != 1:
        raise ValueError(f"Expected one 2026-09 period for IEK {stage1.get('sku')}")
    period = septembers[0]
    raw = period.get("stock")
    quantity = numeric_quantity(raw)
    reference = (period.get("source_cells") or {}).get("stock")
    stock_file = (stage1.get("source_files") or {}).get("stock") or {}
    refs = []
    if isinstance(reference, dict) and isinstance(reference.get("row"), int) \
            and isinstance(reference.get("column"), int):
        refs.append(cell_ref(
            stock_file.get("file", ""), stock_file.get("sheet", ""),
            reference["row"], reference["column"],
        ))
    quality = "historical" if quantity is not None else (
        "missing" if raw is None else "invalid_source"
    )
    if quality == "invalid_source":
        notes.append("invalid_historical_stock_value")
    return evidence(
        quantity, quality, "2026-09-01", "beginning_of_month", refs,
        "Остаток IEK на начало сентября 2026; не является текущим остатком "
        "на дату планирования или файла товаров в пути."
    ), notes


def build_record(stage1: dict, stage1_sha256: str,
                 iek: dict[str, dict], systeme: dict[str, dict]) -> dict:
    if stage1.get("source_commit") != SOURCE_COMMIT:
        raise ValueError(f"Stage-1 source commit mismatch for {stage1.get('brand')} / {stage1.get('sku')}")
    brand, sku = stage1["brand"], stage1["sku"]
    notes = []
    if brand == "IEK":
        historical, historical_notes = iek_historical_stock(stage1)
        notes.extend(historical_notes)
        current = evidence(
            None, "missing", None, "no_current_stock_snapshot", [],
            "Сентябрьский остаток IEK — на начало месяца; движения прихода после него "
            "неполны, поэтому актуальный остаток не выводится."
        )
        transit_source = iek.get(sku)
        if transit_source is None:
            transit = evidence(
                None, "missing", None, "shipment_columns_D_to_I", [],
                "SKU отсутствует в файле пути IEK; отсутствие строки не означает нулевой товар в пути."
            )
            notes.append("sku_missing_from_iek_transit_source")
        else:
            transit = transit_source["evidence"]
            notes.extend(transit_source["notes"])
        extra = {"historical_stock": historical}
    elif brand == "Systeme Electric":
        snapshot = systeme.get(sku)
        if snapshot is None:
            current = evidence(
                None, "missing", None, "reported_total_stock_AX", [],
                "SKU отсутствует в сводке Systeme; текущий остаток неизвестен."
            )
            transit = evidence(
                None, "missing", None, "reported_goods_in_transit_BC", [],
                "SKU отсутствует в сводке Systeme; товар в пути неизвестен."
            )
            notes.append("sku_missing_from_systeme_snapshot_source")
        else:
            current, transit = snapshot["current_stock"], snapshot["goods_in_transit"]
            notes.extend(snapshot["notes"])
        extra = {}
    else:
        raise ValueError(f"Unsupported brand: {brand!r}")
    return {
        "source_commit": SOURCE_COMMIT,
        "stage1_sha256": stage1_sha256,
        "brand": brand,
        "sku": sku,
        "current_stock": current,
        "goods_in_transit": transit,
        **extra,
        "notes": notes,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stage1", type=Path, default=STAGE1_FILE)
    parser.add_argument("--output", type=Path, default=OUTPUT_FILE)
    args = parser.parse_args()
    if not args.stage1.is_file():
        parser.error(f"{args.stage1} is missing; run Stage 1 first")

    stage1_sha256 = file_sha256(args.stage1)
    iek, systeme = load_source_balances()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    seen: set[tuple[str, str]] = set()
    count = 0
    temporary = None
    try:
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=args.output.parent,
                                         prefix="balances-", suffix=".tmp", delete=False) as target:
            temporary = Path(target.name)
            with args.stage1.open(encoding="utf-8") as source:
                for line_number, line in enumerate(source, 1):
                    stage1 = json.loads(line)
                    key = (stage1["brand"], stage1["sku"])
                    if key in seen:
                        raise ValueError(f"Duplicate Stage-1 SKU at line {line_number}: {key}")
                    seen.add(key)
                    result = build_record(stage1, stage1_sha256, iek, systeme)
                    target.write(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n")
                    count += 1
        os.replace(temporary, args.output)
    except (ValueError, KeyError, json.JSONDecodeError):
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise
    print(f"Source balances for {count} Stage-1 SKUs: {args.output}")
    print(f"Stage-1 SHA-256: {stage1_sha256}")


if __name__ == "__main__":
    main()
