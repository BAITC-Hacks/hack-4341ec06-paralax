"""Stage-0 read-only audit of the 12 workbooks held in a Git ref.

Run from repository root. JSON output goes to ignored .analysis/ and contains
source-level diagnostics; it never writes or modifies Excel workbooks.
"""

from __future__ import annotations

import argparse
import io
import json
import re
import subprocess
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def git_bytes(*args: str) -> bytes:
    return subprocess.check_output(["git", *args])


def source_kind(path: str) -> tuple[str, str]:
    brand = "IEK" if path.startswith("data/IEK/") else "Systeme Electric"
    name = path.rsplit("/", 1)[-1].lower()
    if name.startswith("moq"):
        return brand, "moq"
    if name.startswith("динамика продаж"):
        return brand, "documents"
    if name.startswith("ежемесячные остатки"):
        return brand, "monthly_stock"
    if name.startswith("ежемесячные продажи"):
        return brand, "monthly_sales"
    if name.startswith("путь ") or name.startswith("товар в пути"):
        return brand, "transit"
    return brand, "seasonality"


def key_column(brand: str, kind: str) -> tuple[int, int]:
    return {
        "IEK": {
            "moq": (1, 2), "documents": (3, 2), "monthly_stock": (2, 4),
            "monthly_sales": (1, 3), "transit": (0, 2),
        },
        "Systeme Electric": {
            "moq": (2, 2), "documents": (3, 2), "monthly_stock": (2, 2),
            "monthly_sales": (1, 3), "transit": (2, 3),
        },
    }[brand][kind]


def number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
        try:
            return float(text) if text else None
        except ValueError:
            return None
    return None


def parse_day(value: object) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, str):
        match = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", value)
        if match:
            day, month, year = map(int, match.groups())
            try:
                return date(year, month, day).isoformat()
            except ValueError:
                return None
    return None


def profile(ref: str) -> dict:
    paths = [
        raw.decode("utf-8")
        for raw in git_bytes("ls-tree", "-r", "-z", "--name-only", ref, "data").split(b"\0")
        if raw.lower().endswith(b".xlsx")
    ]
    reports: list[dict] = []
    keys: dict[str, dict[str, set[str]]] = defaultdict(dict)
    examples: dict[str, dict[str, dict]] = defaultdict(lambda: defaultdict(dict))
    for path in paths:
        brand, kind = source_kind(path)
        workbook = load_workbook(io.BytesIO(git_bytes("show", f"{ref}:{path}")), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        if kind == "seasonality":
            reports.append({"brand": brand, "kind": kind, "path": path, "sheet": sheet.title,
                            "rows_reported": sheet.max_row, "columns_reported": sheet.max_column,
                            "note": "aggregate money/brand seasonality; no SKU key"})
            workbook.close()
            continue
        key_idx, start = key_column(brand, kind)
        seen: set[str] = set()
        counters: Counter[str] = Counter()
        units: Counter[str] = Counter()
        warehouses: Counter[str] = Counter()
        days: list[str] = []
        duplicates: list[str] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=start, values_only=True), start=start):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            counters["data_rows"] += 1
            raw_key = row[key_idx] if key_idx < len(row) else None
            if raw_key is None or not str(raw_key).strip():
                counters["missing_key"] += 1
                continue
            if not isinstance(raw_key, str):
                counters["nontext_key"] += 1
            key = str(raw_key)
            if key != key.strip():
                counters["key_with_outer_whitespace"] += 1
            if not key.endswith("_"):
                counters["key_without_underscore"] += 1
            if kind != "documents" and key in seen:
                counters["duplicate_key_rows"] += 1
                if len(duplicates) < 5:
                    duplicates.append(key)
            seen.add(key)
            record = examples[brand][key]
            source_rows = record.setdefault("source_rows", {})
            source_rows.setdefault(kind, [])
            if len(source_rows[kind]) < 2:
                source_rows[kind].append(row_number)
            if kind == "documents":
                units[str(row[5] or "")] += 1
                warehouses[str(row[6] or "")] += 1
                day = parse_day(row[0])
                if day:
                    days.append(day)
                else:
                    counters["bad_document_date"] += 1
                quantity = number(row[7])
                if quantity is None:
                    counters["non_numeric_document_quantity"] += 1
                elif quantity < 0:
                    counters["negative_document_quantity"] += 1
                elif quantity > 0:
                    counters["positive_document_quantity"] += 1
                else:
                    counters["zero_document_quantity"] += 1
                record["document_rows"] = record.get("document_rows", 0) + 1
            elif kind in {"monthly_sales", "monthly_stock"}:
                month_start = (3 if brand == "IEK" else 4) if kind == "monthly_stock" else (2 if brand == "IEK" else 4)
                month_end = (36 if brand == "IEK" else 37) if kind == "monthly_stock" else (35 if brand == "IEK" else 37)
                vals = row[month_start:month_end]
                counters["negative_month_values"] += sum(1 for value in vals if (number(value) or 0) < 0)
                counters["blank_month_values"] += sum(1 for value in vals if value is None or value == "")
                if kind == "monthly_stock":
                    units[str(row[1 if brand == "IEK" else 3] or "")] += 1
                    record["latest_monthly_stock"] = row[month_end - 1] if len(row) >= month_end else None
                else:
                    record["latest_monthly_sales"] = row[month_end - 1] if len(row) >= month_end else None
            elif kind == "moq":
                value = number(row[4])
                if value is None or value <= 0:
                    counters["blank_or_nonpositive_rule"] += 1
                record["supplier_rule"] = row[4]
            elif kind == "transit":
                if brand == "IEK":
                    amounts = [number(value) for value in row[3:9]]
                    record["transit_sum"] = sum(value or 0 for value in amounts)
                    counters["negative_transit_cells"] += sum(1 for value in amounts if value is not None and value < 0)
                else:
                    record["transit_value"] = row[54] if len(row) > 54 else None
                    record["snapshot_stock"] = row[49] if len(row) > 49 else None
                    record["free_stock"] = row[51] if len(row) > 51 else None
            record[f"in_{kind}"] = True
        keys[brand][kind] = seen
        reports.append({"brand": brand, "kind": kind, "path": path, "sheet": sheet.title,
                        "rows_reported": sheet.max_row, "columns_reported": sheet.max_column,
                        "unique_keys": len(seen), "counters": dict(counters),
                        "units_top": units.most_common(8), "warehouses_top": warehouses.most_common(8),
                        "date_min": min(days) if days else None, "date_max": max(days) if days else None,
                        "duplicate_key_examples": duplicates})
        workbook.close()
    coverage = {}
    for brand, collections in keys.items():
        shared = set.intersection(*collections.values())
        coverage[brand] = {
            "all_five_sources": len(shared),
            "source_unique_keys": {kind: len(values) for kind, values in collections.items()},
            "missing_from_monthly_sales": len(collections["monthly_stock"] - collections["monthly_sales"]),
            "missing_from_monthly_stock": len(collections["monthly_sales"] - collections["monthly_stock"]),
            "missing_from_transit": len(collections["monthly_sales"] - collections["transit"]),
            "missing_from_moq": len(collections["monthly_sales"] - collections["moq"]),
            "examples": [
                {"sku": sku, **examples[brand][sku]}
                for sku in sorted(shared)
                if examples[brand][sku].get("document_rows", 0) >= 2
            ][:5],
        }
    return {
        "ref": ref,
        "ref_commit": git_bytes("rev-parse", ref).decode("ascii").strip(),
        "workbooks": len(paths),
        "reports": reports,
        "coverage": coverage,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ref", default="origin/main")
    parser.add_argument("--output", type=Path, default=Path(".analysis/source-audit.json"))
    args = parser.parse_args()
    if args.ref != "origin/main":
        raise SystemExit("Stage 0 reads origin/main only; review other refs explicitly")
    result = profile(args.ref)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"Audited {result['workbooks']} workbooks. Local report: {args.output}")


if __name__ == "__main__":
    main()
