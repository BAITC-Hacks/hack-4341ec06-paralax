"""Read Excel workbooks from a Git ref without changing the checked-out branch.

The JSON output is local analysis material. It may contain source values and
must not be committed or sent to an external service without review.
"""

from __future__ import annotations

import argparse
import io
import json
import subprocess
from pathlib import Path

from openpyxl import load_workbook


def git_bytes(*args: str) -> bytes:
    return subprocess.check_output(["git", *args])


def summarize(ref: str, preview_rows: int = 12, preview_cols: int = 100) -> list[dict]:
    paths = [
        raw.decode("utf-8")
        for raw in git_bytes("ls-tree", "-r", "-z", "--name-only", ref, "data").split(b"\0")
        if raw.lower().endswith(b".xlsx")
    ]
    result: list[dict] = []
    for path in paths:
        source = git_bytes("show", f"{ref}:{path}")
        workbook = load_workbook(io.BytesIO(source), read_only=True, data_only=True)
        sheets = []
        for sheet in workbook.worksheets:
            preview = []
            for row in sheet.iter_rows(min_row=1, max_row=preview_rows, max_col=preview_cols, values_only=True):
                values = [str(value).strip()[:100] if value is not None else "" for value in row]
                if any(values):
                    preview.append(values)
            sheets.append({
                "name": sheet.title,
                "rows_reported": sheet.max_row,
                "columns_reported": sheet.max_column,
                "preview": preview,
            })
        workbook.close()
        result.append({"path": path, "bytes": len(source), "sheets": sheets})
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ref", default="origin/main", help="read-only Git ref with source data")
    parser.add_argument("--output", type=Path, default=Path(".analysis/workbooks.json"))
    args = parser.parse_args()
    if args.ref != "origin/main":
        raise SystemExit("Stage 0 only reads origin/main by default; review other refs explicitly")
    result = summarize(args.ref)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"Read {len(result)} workbooks and {sum(len(item['sheets']) for item in result)} sheets; wrote {args.output}")


if __name__ == "__main__":
    main()
