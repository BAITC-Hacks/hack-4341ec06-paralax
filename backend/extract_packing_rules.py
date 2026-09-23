"""Stage 6: extract supplier order rules from the immutable source workbooks.

IEK's "Мин. разр. к отгр." is a minimum dispatch quantity, while Systeme
Electric's "Кратность" is a shipment multiple. They must not be interchanged.
Derived partner data is written only to the ignored .analysis directory.
"""

from __future__ import annotations

import argparse
import io
import json
import math
import os
import subprocess
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from clean_demand import SOURCE_COMMIT

OUTPUT_FILE = Path(".analysis/packing-rules/source-rules.jsonl")
RULE_BOOKS = {
    "IEK": {
        "file": "data/IEK/MOQ  ИЭК.xlsx",
        "sheet": "Лист7",
        "sku_column": 2,
        "rule_column": 5,
        "rule_kind": "minimum_dispatch",
        "sku_header": "Код 1с",
        "rule_header": "Мин. разр. к отгр.",
        "first_row": 2,
    },
    "Systeme Electric": {
        "file": "data/Systeme electric/MOQ SystemElectric.xlsx",
        "sheet": "Лист_1",
        "sku_column": 3,
        "rule_column": 5,
        "rule_kind": "multiple",
        "sku_header": "Номенклатура.Код",
        "rule_header": "Кратность",
        "first_row": 3,
    },
}


def positive_integer(value: object) -> int | None:
    """Accept only an explicit, finite positive integer Excel number."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if not math.isfinite(value) or value <= 0 or not float(value).is_integer():
        return None
    return int(value)


def source_value(value: object) -> str | int | float | None:
    """Keep the visible source value while producing standards-compliant JSON."""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else str(value)
    return None


def parse_rule_rows(
    brand: str, numbered_rows: Iterable[tuple[int, Sequence[object]]]
) -> dict[str, dict]:
    """Group source cells by 1C SKU, retaining invalid and duplicate evidence."""
    config = RULE_BOOKS[brand]
    sku_index = config["sku_column"] - 1
    rule_index = config["rule_column"] - 1
    grouped: dict[str, list[tuple[int, object]]] = defaultdict(list)
    for row_number, row in numbered_rows:
        raw_sku = row[sku_index] if len(row) > sku_index else None
        if raw_sku is None or not str(raw_sku).strip():
            continue  # Blank separator, title, or summary without a 1C SKU.
        sku = str(raw_sku).strip()
        value = row[rule_index] if len(row) > rule_index else None
        grouped[sku].append((row_number, value))

    result = {}
    for sku, entries in grouped.items():
        values = [positive_integer(value) for _, value in entries]
        source_refs = [
            {
                "file": config["file"],
                "sheet": config["sheet"],
                "row": row_number,
                "column": config["rule_column"],
                "cell": f"{get_column_letter(config['rule_column'])}{row_number}",
                "source_value": source_value(value),
            }
            for row_number, value in entries
        ]
        notes = []
        if len(entries) > 1 and all(value is not None and value == values[0] for value in values):
            status, quantity = "valid", values[0]
            notes.append("duplicate_identical_rule")
        elif len(entries) > 1:
            status, quantity = "conflict", None
            notes.append("duplicate_rule_conflict")
        elif values[0] is not None:
            status, quantity = "valid", values[0]
        elif entries[0][1] is None or isinstance(entries[0][1], str) and not entries[0][1].strip():
            status, quantity = "missing", None
            notes.append("blank_rule_cell")
        else:
            status, quantity = "invalid", None
            notes.append("rule_must_be_positive_integer_number")
        result[sku] = {
            "source_commit": SOURCE_COMMIT,
            "brand": brand,
            "sku": sku,
            "rule_kind": config["rule_kind"],
            "quantity": quantity,
            "status": status,
            "source_refs": source_refs,
            "notes": notes,
        }
    return result


def load_source_rules() -> list[dict]:
    records = []
    for brand, config in RULE_BOOKS.items():
        source = subprocess.check_output(["git", "show", f"{SOURCE_COMMIT}:{config['file']}"])
        workbook = load_workbook(io.BytesIO(source), read_only=True, data_only=True)
        try:
            sheet = workbook[config["sheet"]]
            header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            if header[config["sku_column"] - 1] != config["sku_header"] \
                    or header[config["rule_column"] - 1] != config["rule_header"]:
                raise ValueError(f"Unexpected {brand} order-rule headers")
            parsed = parse_rule_rows(
                brand,
                enumerate(sheet.iter_rows(min_row=config["first_row"], values_only=True),
                          start=config["first_row"]),
            )
            records.extend(parsed.values())
        finally:
            workbook.close()
    return records


def write_rules(output: Path, records: list[dict]) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(
            "w", encoding="utf-8", dir=output.parent,
            prefix="packing-rules-", suffix=".tmp", delete=False,
        ) as target:
            temporary = Path(target.name)
            for record in records:
                target.write(json.dumps(record, ensure_ascii=False, allow_nan=False,
                                        separators=(",", ":")) + "\n")
        os.replace(temporary, output)
    except (OSError, TypeError, ValueError):
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        raise


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT_FILE)
    args = parser.parse_args()
    records = load_source_rules()
    write_rules(args.output, records)
    counts = Counter(record["status"] for record in records)
    print(f"Source order rules for {len(records)} SKUs: {args.output}")
    print("Rule statuses: " + ", ".join(f"{key}={counts[key]}" for key in sorted(counts)))


if __name__ == "__main__":
    main()
