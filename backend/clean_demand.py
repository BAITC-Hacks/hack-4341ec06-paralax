"""Stage 1: traceable monthly clean demand from the source Excel workbooks.

Reads the immutable source Git commit; writes derived partner data only to the
ignored .analysis directory. No API contract or raw workbook is modified.
"""

from __future__ import annotations

import argparse
import csv
import html
import io
import json
import math
import re
import statistics
import subprocess
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

SOURCE_COMMIT = "a7a1d3bcd25ab137eac7915f29a4e5c2f5a0f5bc"
MONTHS = [f"{year}-{month:02d}" for year in (2024, 2025, 2026)
          for month in range(1, 13) if year < 2026 or month <= 9]
SOURCES = {
    "IEK": {
        "documents": "data/IEK/Динамика продаж_2025-2026.xlsx",
        "sales": "data/IEK/Ежемесячные продажи в количественном выражении за последние 2 года.xlsx",
        "stock": "data/IEK/Ежемесячные остатки продукции за последние 2 года  ИЭК.xlsx",
    },
    "Systeme Electric": {
        "documents": "data/Systeme electric/Динамика продаж_Syseme Electric_2025-2026.xlsx",
        "sales": "data/Systeme electric/Ежемесячные продажи в кол-м выражении SystemElectric 2024-2026.xlsx",
        "stock": "data/Systeme electric/Ежемесячные остатки SystemElectric 2024-2026.xlsx",
    },
}
SHEET = {"documents": "Лист_1", "sales": "Лист_1", "stock": "Лист_1"}


def number(value: object) -> float | None:
    if isinstance(value, bool) or value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, str):
        try:
            result = float(value.strip().replace(" ", "").replace("\u00a0", "").replace(",", "."))
            return result if math.isfinite(result) else None
        except ValueError:
            return None
    return None


def doc_month(value: object) -> str | None:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m")
    if isinstance(value, str):
        match = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", value)
        if match:
            day, month, year = map(int, match.groups())
            try:
                date(year, month, day)
                return f"{year}-{month:02d}"
            except ValueError:
                pass
    return None


def sheet_from_git(path: str):
    source = subprocess.check_output(["git", "show", f"{SOURCE_COMMIT}:{path}"])
    book = load_workbook(io.BytesIO(source), read_only=True, data_only=True)
    return book, book["Лист_1"]


def load_sources() -> dict:
    products: dict[tuple[str, str], dict] = {}
    for brand, paths in SOURCES.items():
        book, sheet = sheet_from_git(paths["sales"])
        start = 3
        first_month_col = 3 if brand == "IEK" else 5
        for row_no, row in enumerate(sheet.iter_rows(min_row=start, values_only=True), start):
            if len(row) < 2 or row[1] is None or not str(row[1]).strip():
                continue
            sku = str(row[1]).strip()
            product = products.setdefault((brand, sku), {"brand": brand, "sku": sku, "months": {}, "docs": []})
            product["name"] = str(row[0] or "")
            for index, month in enumerate(MONTHS):
                col = first_month_col + index
                raw = row[col - 1] if len(row) >= col else None
                product["months"].setdefault(month, {})["sales"] = {
                    "value": number(raw), "raw_cell": str(raw) if raw is not None else None,
                    "ref": {"row": row_no, "column": col},
                }
        book.close()

        book, sheet = sheet_from_git(paths["stock"])
        first_row = 4 if brand == "IEK" else 2
        first_month_col = 4 if brand == "IEK" else 5
        for row_no, row in enumerate(sheet.iter_rows(min_row=first_row, values_only=True), first_row):
            if len(row) < 3 or row[2] is None or not str(row[2]).strip():
                continue
            sku = str(row[2]).strip()
            product = products.setdefault((brand, sku), {"brand": brand, "sku": sku, "months": {}, "docs": []})
            product.setdefault("name", str(row[0 if brand == "IEK" else 1] or ""))
            product["unit"] = str(row[1 if brand == "IEK" else 3] or "")
            for index, month in enumerate(MONTHS):
                col = first_month_col + index
                raw = row[col - 1] if len(row) >= col else None
                product["months"].setdefault(month, {})["stock"] = {
                    "value": number(raw), "raw_cell": str(raw) if raw is not None else None,
                    "ref": {"row": row_no, "column": col},
                }
        book.close()

        book, sheet = sheet_from_git(paths["documents"])
        for row_no, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if len(row) < 8 or row[3] is None or not str(row[3]).strip():
                continue
            sku = str(row[3]).strip()
            product = products.setdefault((brand, sku), {"brand": brand, "sku": sku, "months": {}, "docs": []})
            product.setdefault("name", str(row[4] or ""))
            month = doc_month(row[0])
            quantity = number(row[7])
            if month not in MONTHS or quantity is None:
                continue
            product["docs"].append({
                "month": month, "row": row_no, "document": str(row[1] or ""),
                "quantity": quantity, "unit": str(row[5] or ""),
            })
        book.close()
    return products


def percentile(values: list[float], fraction: float) -> float:
    ordered = sorted(values)
    position = (len(ordered) - 1) * fraction
    low, high = math.floor(position), math.ceil(position)
    return ordered[low] + (ordered[high] - ordered[low]) * (position - low)


def one_off_documents(documents: list[dict], unit: str) -> tuple[dict[str, list[dict]], dict]:
    positives = [item["quantity"] for item in documents
                 if item["quantity"] > 0 and (not unit or item["unit"] == unit)]
    if len(positives) < 8:
        return {}, {"eligible_document_count": len(positives), "threshold": None}
    median = statistics.median(positives)
    mad = statistics.median(abs(value - median) for value in positives)
    q1, q3 = percentile(positives, 0.25), percentile(positives, 0.75)
    threshold = max(4 * median, median + 6 * 1.4826 * mad, q3 + 3 * (q3 - q1), median + 5)
    by_month: dict[str, list[dict]] = defaultdict(list)
    for item in documents:
        if item["quantity"] > threshold and (not unit or item["unit"] == unit):
            by_month[item["month"]].append({
                "row": item["row"], "document": item["document"],
                "quantity": item["quantity"], "excluded_excess": item["quantity"] - median,
            })
    return by_month, {
        "eligible_document_count": len(positives), "median": median, "mad": mad,
        "q1": q1, "q3": q3, "threshold": threshold,
    }


def normal_reference(periods: list[dict], index: int) -> tuple[float | None, float, list[str]]:
    """Nearby available-stock months; optional SKU-specific same-month seasonality."""
    candidates = [
        (abs(j - index), j, period["adjusted_sales"])
        for j, period in enumerate(periods)
        if j != index and abs(j - index) <= 6
        and period["stock"] is not None and period["stock"] > 0
        and period["adjusted_sales"] is not None
        and period["adjusted_sales"] > 0 and not period["partial_period"]
    ]
    candidates.sort()
    nearest = candidates[:6]
    if len(nearest) < 2:
        return None, 1.0, []
    baseline = statistics.median(value for _, _, value in nearest)
    eligible = [p for p in periods if p["stock"] is not None and p["stock"] > 0
                and p["adjusted_sales"] is not None and p["adjusted_sales"] > 0
                and not p["partial_period"]]
    same_month = [p["adjusted_sales"] for p in eligible
                  if p["period"][-2:] == periods[index]["period"][-2:]]
    multiplier = 1.0
    if len(eligible) >= 12 and len(same_month) >= 2:
        overall = statistics.median(p["adjusted_sales"] for p in eligible)
        if overall > 0:
            multiplier = min(1.5, max(0.5, statistics.median(same_month) / overall))
    return baseline * multiplier, multiplier, [periods[j]["period"] for _, j, _ in nearest]


def clean_product(product: dict) -> dict:
    brand, sku = product["brand"], product["sku"]
    outliers, outlier_rule = one_off_documents(product["docs"], product.get("unit", ""))
    document_totals: dict[str, float] = defaultdict(float)
    for document in product["docs"]:
        if not product.get("unit") or document["unit"] == product["unit"]:
            document_totals[document["month"]] += document["quantity"]
    periods = []
    for month in MONTHS:
        source = product["months"].get(month, {})
        sales = source.get("sales", {})
        stock = source.get("stock", {})
        raw_sales = sales.get("value")
        observed = max(0.0, raw_sales) if raw_sales is not None else None
        flagged = outliers.get(month, [])
        proposed_exclusion = sum(item["excluded_excess"] for item in flagged)
        document_total = document_totals.get(month)
        reconciled = (raw_sales is not None and document_total is not None
                      and abs(document_total - raw_sales) <= max(2.0, 0.05 * max(abs(raw_sales), abs(document_total))))
        excluded = min(observed or 0, proposed_exclusion) if reconciled else 0.0
        adjusted = max(0.0, observed - excluded) if observed is not None else None
        reasons = []
        if raw_sales is None:
            reasons.append("missing_sales")
        elif raw_sales < 0:
            reasons.append("negative_net_sales_return")
        if flagged:
            reasons.append("one_off_document_median_mad_iqr" if reconciled
                           else "one_off_document_unreconciled_no_exclusion")
        if reconciled and proposed_exclusion > (observed or 0):
            reasons.append("document_month_mismatch_exclusion_capped")
        partial = month == "2026-09"
        if partial:
            reasons.append("partial_period")
        periods.append({
            "period": month, "raw_sales": raw_sales, "observed_demand": observed,
            "excluded_one_off": excluded, "adjusted_sales": adjusted,
            "document_month_total": document_total,
            "document_reconciled": reconciled if flagged else None,
            "stock": stock.get("value"), "estimated_lost_demand": 0.0,
            "clean_demand": adjusted, "reasons": reasons,
            "partial_period": partial, "seasonality_multiplier": 1.0,
            "reference_periods": [], "outlier_documents": flagged,
            "source_cells": {"sales": sales.get("ref"), "stock": stock.get("ref")},
        })
    for index, period in enumerate(periods):
        if period["partial_period"] or period["stock"] != 0 or period["excluded_one_off"] > 0:
            continue
        expected, multiplier, reference_months = normal_reference(periods, index)
        if expected is None:
            period["reasons"].append("stock_zero_insufficient_reference")
            continue
        period["seasonality_multiplier"] = multiplier
        period["reference_periods"] = reference_months
        adjusted = period["adjusted_sales"]
        if adjusted is None or adjusted <= 0.25 * expected:
            period["estimated_lost_demand"] = max(0.0, expected - (adjusted or 0.0))
            period["clean_demand"] = (adjusted or 0.0) + period["estimated_lost_demand"]
            period["reasons"].append("possible_stockout_beginning_stock_zero" if brand == "IEK"
                                      else "possible_stockout_stock_snapshot_zero")
            if multiplier != 1.0:
                period["reasons"].append("sku_calendar_month_seasonality")
    return {
        "source_commit": SOURCE_COMMIT, "brand": brand, "sku": sku,
        "name": product.get("name", ""), "unit": product.get("unit"),
        "source_files": {kind: {"file": path, "sheet": SHEET[kind]}
                         for kind, path in SOURCES[brand].items()},
        "outlier_rule": outlier_rule, "document_movements": product["docs"],
        "periods": periods,
    }


def render_html(result: dict, path: Path) -> None:
    periods = result["periods"]
    values = [value for p in periods for value in (p["raw_sales"], p["clean_demand"])
              if value is not None and value >= 0]
    maximum = max(values or [1]) or 1
    width, height = 1120, 350
    def lines(field: str, color: str) -> str:
        coordinates = [
            (30 + i * 32, height - 30 - max(0, p[field]) / maximum * (height - 65))
            if p[field] is not None else None for i, p in enumerate(periods)
        ]
        segments = [f'<line x1="{left[0]}" y1="{left[1]:.1f}" x2="{right[0]}" '
                    f'y2="{right[1]:.1f}" stroke="{color}" stroke-width="2"/>'
                    for left, right in zip(coordinates, coordinates[1:]) if left and right]
        dots = [f'<circle cx="{point[0]}" cy="{point[1]:.1f}" r="2.5" fill="{color}"/>'
                for point in coordinates if point]
        return "".join(segments + dots)
    rows = []
    for p in periods:
        def show(value: object) -> str:
            return "—" if value is None else html.escape(str(round(value, 2) if isinstance(value, float) else value))
        documents = ", ".join(f"#{item['row']}: {show(item['quantity'])}"
                              for item in p["outlier_documents"])
        if documents:
            documents += f" (Σ документов {show(p['document_month_total'])})"
        rows.append("<tr>" + "".join(f"<td>{show(p[key])}</td>" for key in
                    ("period", "raw_sales", "excluded_one_off", "estimated_lost_demand", "clean_demand", "stock"))
                    + f"<td>{html.escape(', '.join(p['reasons']))}</td><td>{documents}</td></tr>")
    title = html.escape(f"{result['brand']} / {result['sku']}")
    markup = f"""<!doctype html><html lang="ru"><meta charset="utf-8"><title>Чистый спрос — {title}</title>
<style>body{{font:15px system-ui;margin:32px;max-width:1250px}}table{{border-collapse:collapse;width:100%}}
td,th{{padding:7px;border-bottom:1px solid #ddd;text-align:left}}th{{background:#f3f5f7}}
svg{{width:100%;height:auto}}.legend span{{margin-right:24px}}</style>
<h1>Чистый спрос: {title}</h1><p>Исходный Git-коммит: <code>{SOURCE_COMMIT}</code>.
Синим — raw sales, оранжевым — clean demand. Сентябрь 2026 — неполный месяц.</p>
<p>Порог крупного документа: {html.escape(str(result['outlier_rule'].get('threshold') or 'недостаточно строк'))};
медиана: {html.escape(str(result['outlier_rule'].get('median', '—')))}. Номера строк документов показаны в таблице.</p>
<div class="legend"><span style="color:#2563eb">● raw sales</span><span style="color:#e86d24">● clean demand</span></div>
<svg viewBox="0 0 {width} {height}" role="img" aria-label="Сравнение исходных продаж и чистого спроса">
<line x1="30" y1="{height-30}" x2="{width-20}" y2="{height-30}" stroke="#999"/>
{lines('raw_sales', '#2563eb')}{lines('clean_demand', '#e86d24')}</svg>
<table><thead><tr><th>Период</th><th>Raw sales</th><th>Исключено</th><th>Восстановлено</th>
<th>Clean demand</th><th>Остаток</th><th>Причина</th><th>Помеченные документы</th></tr></thead><tbody>{''.join(rows)}</tbody></table>
<p>Полная трассировка колонок и строк источника, а также документов-выбросов находится в JSONL для этого SKU.</p></html>"""
    path.write_text(markup, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sku", action="append", help="also write HTML/CSV/JSON for this SKU; repeatable")
    parser.add_argument("--brand", choices=list(SOURCES), help="required if SKU exists for both brands")
    parser.add_argument("--output-dir", type=Path, default=Path(".analysis/clean-demand"))
    args = parser.parse_args()
    if args.brand and not args.sku:
        parser.error("--brand requires --sku")
    if subprocess.check_output(["git", "rev-parse", SOURCE_COMMIT]).decode().strip() != SOURCE_COMMIT:
        raise SystemExit("Source Git commit is unavailable")
    products = load_sources()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    selected = []
    output = args.output_dir / "all-skus.jsonl"
    with output.open("w", encoding="utf-8") as stream:
        for key in sorted(products):
            result = clean_product(products[key])
            stream.write(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n")
            if args.sku and key[1] in args.sku and (not args.brand or args.brand == key[0]):
                selected.append(result)
    for result in selected:
        stem = re.sub(r"[^A-Za-z0-9_-]", "-", result["sku"])
        stem = ("iek-" if result["brand"] == "IEK" else "systeme-") + stem
        (args.output_dir / f"{stem}.json").write_text(
            json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        with (args.output_dir / f"{stem}.csv").open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=[
                "period", "raw_sales", "excluded_one_off", "estimated_lost_demand",
                "clean_demand", "stock", "reasons"])
            writer.writeheader()
            for period in result["periods"]:
                writer.writerow({key: ",".join(period[key]) if key == "reasons" else period[key]
                                 for key in writer.fieldnames})
        render_html(result, args.output_dir / f"{stem}.html")
    missing = set(args.sku or []) - {result["sku"] for result in selected}
    if missing:
        raise SystemExit(f"SKUs not found: {', '.join(sorted(missing))}")
    print(f"Clean demand for {len(products)} SKUs: {output}")
    for result in selected:
        print(f"Trace: {args.output_dir / (('iek-' if result['brand'] == 'IEK' else 'systeme-') + re.sub(r'[^A-Za-z0-9_-]', '-', result['sku']) + '.html')}")


if __name__ == "__main__":
    main()
